"""Opt-in Excel-for-Windows native workbook validation and macro runner.

This module is CLI-only. It is intentionally NOT connected to FastAPI endpoints:
a web upload must never be able to trigger Microsoft Excel/VBA execution.

Native Excel is not a sandbox. A macro can access the operating system,
network, and other files despite copying the workbook before execution.
Use macro mode only on reviewed workbooks in a disposable Windows VM.
"""

import argparse
from datetime import datetime, timezone
from html import escape
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sys
from uuid import uuid4

from openpyxl import load_workbook
from webapp.workbook import validate_upload

MAX_FORMULAS = 20_000
MAX_CHANGED_DETAILS = 2_000
MACRO_PATTERN = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*(?:\.[A-Za-z_][A-Za-z0-9_]*)?$")
TRUST_PHRASE = "I UNDERSTAND VBA CAN RUN SYSTEM COMMANDS"


class NativeExcelError(Exception):
    """Human-readable validation or Excel automation error."""


def _hash_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as data:
        for chunk in iter(lambda: data.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sanitize(value):
    if value is None or isinstance(value, (bool, int, float, str)):
        return value if not isinstance(value, str) else value[:300]
    return str(value)[:300]


def _collect_formulas(path):
    """Snapshot the on-disk cached value for each formula before Excel opens it."""
    source = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    cached = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    formulas = []
    try:
        for sheet in source:
            saved_values = cached[sheet.title]
            for cell in sheet._cells.values():
                if cell.data_type != "f":
                    continue
                if len(formulas) >= MAX_FORMULAS:
                    raise NativeExcelError(
                        f"Workbook contains more than {MAX_FORMULAS} formulas; "
                        "native preview limit exceeded."
                    )
                formulas.append({
                    "sheet": sheet.title,
                    "cell": cell.coordinate,
                    "formula": str(cell.value)[:300],
                    "before": _sanitize(saved_values[cell.coordinate].value),
                })
        return formulas
    finally:
        source.close()
        cached.close()


def _render_html(report):
    """Small, offline, escaped HTML summary; never embeds macro source."""
    meta = [
        ("Execution mode", report["mode"]),
        ("Started (UTC)", report["started_utc"]),
        ("Source", report["source_filename"]),
        ("Excel installed/automation", report["excel_automation"]),
        ("Formulas inventoried", str(report["formulas_inventory_count"])),
        ("Formula results changed", str(report["formula_results_changed"])),
        ("Macro executed", str(report["macro_executed"])),
        ("Source unchanged", str(report.get("original_unchanged"))),
    ]
    bits = ['<!doctype html><html lang="en"><head><meta charset="utf-8">',
            "<title>Data Quality Detective — Excel native audit</title>",
            "<style>body{font:15px system-ui;margin:32px auto;max-width:1100px;"
            "color:#223047}table{border-collapse:collapse;width:100%}"
            "th,td{padding:9px;border:1px solid #d8e1ed;text-align:left}"
            "th{background:#edf3ff}p{line-height:1.6}</style></head><body>",
            "<h1>Excel-native audit</h1>"]
    for label, value in meta:
        bits.append("<p><strong>" + escape(label) + ":</strong> " +
                    escape(str(value)) + "</p>")
    bits.append("<p><strong>Safety:</strong> " +
                escape(report["limitations"]) + "</p>")
    if report["warnings"]:
        bits.append("<h2>Warnings</h2><ul>")
        bits.extend("<li>" + escape(str(warning)) + "</li>"
                    for warning in report["warnings"])
        bits.append("</ul>")
    bits.append("<h2>Formula result changes</h2>")
    bits.append("<table><thead><tr><th>Worksheet</th><th>Cell</th>"
                "<th>Formula</th><th>Cached before</th><th>Excel after</th>"
                "</tr></thead><tbody>")
    for item in report["changed_details"]:
        bits.append("<tr>" + "".join(
            "<td>" + escape(str(item[key])) + "</td>"
            for key in ("sheet", "cell", "formula", "before", "after")
        ) + "</tr>")
    bits.append("</tbody></table></body></html>")
    return "".join(bits)


def _prepare_workbook(source_file, output_dir):
    source = Path(source_file).expanduser().resolve(strict=True)
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise NativeExcelError("Native Excel audit currently supports XLSX/XLSM only.")
    data = source.read_bytes()
    validate_upload(source.name, data)
    output = Path(output_dir).expanduser().resolve()
    if output == source.parent:
        raise NativeExcelError(
            "Choose a separate output directory, not the source workbook folder."
        )
    output.mkdir(parents=True, exist_ok=True)
    run_dir = output / ("excel_native_" + uuid4().hex[:12])
    run_dir.mkdir(parents=False, exist_ok=False)
    workbook_copy = run_dir / source.name
    shutil.copy2(source, workbook_copy)
    return source, run_dir, workbook_copy


def run_native_audit(
    source_file, output_dir, mode="recalculate", macro_name=None,
    confirmation=None, dispatch=None
):
    """Run Excel locally on a temporary copy with a dedicated Excel instance.

    Native mode opens workbooks with macros force-disabled.
    Macro mode requests *Excel's normal Trust Center UI* and requires multiple
    explicit user approvals. It is not sandboxing or guaranteed safe.
    """
    if mode not in {"recalculate", "macro"}:
        raise NativeExcelError("Mode must be recalculate or macro.")
    if mode == "macro":
        if not macro_name or not MACRO_PATTERN.fullmatch(macro_name):
            raise NativeExcelError(
                "Macro name must be an ordinary VBA name such as Module1.RunAudit."
            )
        if confirmation != TRUST_PHRASE:
            raise NativeExcelError(
                "Macro mode requires the exact --confirm-trust phrase. "
                "Review the VBA and run inside a disposable Windows VM."
            )
    elif macro_name is not None:
        raise NativeExcelError("Macro names are only accepted in macro mode.")
    if dispatch is None and sys.platform != "win32":
        raise NativeExcelError(
            "Native Excel automation requires Microsoft Excel on Windows."
        )

    source, run_dir, workbook_copy = _prepare_workbook(source_file, output_dir)
    formulas = _collect_formulas(workbook_copy)
    source_sha256 = _hash_file(source)
    copy_sha256_before = _hash_file(workbook_copy)
    report = {
        "mode": mode,
        "source_filename": source.name,
        "working_copy": str(workbook_copy),
        "source_sha256": source_sha256,
        "copy_sha256_before": copy_sha256_before,
        "copy_sha256_after": None,
        "original_unchanged": None,
        "started_utc": datetime.now(timezone.utc).isoformat(),
        "excel_automation": "pending",
        "formulas_inventory_count": len(formulas),
        "formula_results_changed": 0,
        "changed_details": [],
        "macro_executed": False,
        "macro_name": macro_name if mode == "macro" else None,
        "macro_return_value": None,
        "warnings": [],
        "limitations": (
            "Runs in real Microsoft Excel, not an isolated sandbox. "
            "The source file is never opened for editing, but VBA/add-ins "
            "may still access the operating system, other files, and network. "
            "Excel Trust Center policy is never modified by this tool."
        ),
    }
    excel = None
    workbook = None
    did_initialize = False
    try:
        if dispatch is None:
            import pythoncom
            from win32com.client import DispatchEx
            pythoncom.CoInitialize()
            did_initialize = True
            dispatch = DispatchEx

        excel = dispatch("Excel.Application")
        # ForceDisable=3 means files opened via automation never run VBA.
        # ByUI=2 follows the user's existing Trust Center policy (macro mode).
        excel.AutomationSecurity = 3 if mode == "recalculate" else 2
        excel.EnableEvents = False
        excel.AskToUpdateLinks = False
        excel.DisplayAlerts = (mode == "macro")
        excel.Visible = (mode == "macro")
        workbook = excel.Workbooks.Open(
            str(workbook_copy), UpdateLinks=0, ReadOnly=False,
            IgnoreReadOnlyRecommended=True
        )
        report["excel_automation"] = "opened"
        if mode == "recalculate":
            excel.CalculateFullRebuild()
            report["warnings"].append(
                "Excel calculation reflects installed functions/add-ins and may "
                "differ from the source's saved calculation results."
            )
        else:
            # Do not programmatically enable macros, alter trusted locations,
            # or bypass Trust Center. If macros are blocked, Run will fail.
            # Workbook_Open/other events are disabled, but VBA itself may call
            # anything available to the user's Windows account.
            qualified = "'" + workbook_copy.name.replace("'", "''") + "'!" + macro_name
            result = excel.Run(qualified)
            report["macro_executed"] = True
            report["macro_return_value"] = _sanitize(result)
            report["warnings"].append(
                "VBA ran with the current Windows user's privileges; "
                "working-copy isolation does NOT restrict external side effects."
            )
        details = []
        count = 0
        for formula in formulas:
            new_val = _sanitize(
                workbook.Worksheets(formula["sheet"]).Range(formula["cell"]).Value2
            )
            if formula["before"] != new_val:
                count += 1
                if len(details) < MAX_CHANGED_DETAILS:
                    details.append({
                        "sheet": formula["sheet"],
                        "cell": formula["cell"],
                        "formula": formula["formula"],
                        "before": formula["before"],
                        "after": new_val,
                    })
        report["formula_results_changed"] = count
        report["changed_details"] = details
        if count > MAX_CHANGED_DETAILS:
            report["warnings"].append(
                f"Only the first {MAX_CHANGED_DETAILS} changed formula cells are listed."
            )
        workbook.Save()
        report["excel_automation"] = "completed"
    except Exception as exc:
        report["excel_automation"] = "failed"
        report["warnings"].append(
            "Excel native operation failed or was blocked. No success is claimed. "
            "Reason: " + str(exc)[:400]
        )
    finally:
        try:
            if workbook is not None:
                workbook.Close(SaveChanges=False)
        except Exception:
            report["warnings"].append("Could not confirm Excel workbook closed.")
        try:
            if excel is not None:
                excel.Quit()
        except Exception:
            report["warnings"].append("Could not confirm dedicated Excel instance exited.")
        if did_initialize:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass

    # Compare hashes after Excel has finished. Macros can still write to
    # arbitrary paths: a copy alone does not protect other files.
    report["original_unchanged"] = (_hash_file(source) == source_sha256)
    if not report["original_unchanged"]:
        report["warnings"].append(
            "WARNING: Original input file changed during the native audit. "
            "Inspect immediately; a macro/add-in may have modified external files."
        )
        report["excel_automation"] = "failed"
    if workbook_copy.exists():
        report["copy_sha256_after"] = _hash_file(workbook_copy)

    (run_dir / "native_report.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8"
    )
    (run_dir / "native_report.html").write_text(_render_html(report), encoding="utf-8")
    report["report_folder"] = str(run_dir)
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Run native Excel calculations or expressly approved VBA on a COPY."
    )
    parser.add_argument("--file", required=True, help="Source XLSX/XLSM workbook")
    parser.add_argument("--output-dir", required=True, help="Separate output folder")
    parser.add_argument("--mode", choices=["recalculate", "macro"],
                        default="recalculate")
    parser.add_argument("--macro-name", help="Explicit VBA procedure name, e.g. Module1.Audit")
    parser.add_argument("--confirm-trust", help="Exact high-risk macro acknowledgement")
    args = parser.parse_args(argv)
    try:
        result = run_native_audit(
            args.file, args.output_dir, mode=args.mode,
            macro_name=args.macro_name, confirmation=args.confirm_trust
        )
    except (NativeExcelError, FileNotFoundError, OSError) as exc:
        parser.exit(2, f"Excel native audit: {exc}\n")
    print(json.dumps({
        "status": result["excel_automation"],
        "report_folder": result["report_folder"],
        "formula_results_changed": result["formula_results_changed"],
        "macro_executed": result["macro_executed"],
        "warnings": result["warnings"],
    }, indent=2))
    return 0 if result["excel_automation"] == "completed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
