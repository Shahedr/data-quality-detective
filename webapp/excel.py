"""Read-only inspection of workbook list/dropdown validation metadata.

Never evaluates VBA, workbook formulas, macros, external connections, or links.
Supports literal comma-separated lists and ordinary one-column Excel ranges.
All other Excel validation formulas remain explicitly unverified.
"""

from decimal import Decimal, InvalidOperation
import io
import re
from zipfile import ZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

# Allows simple Excel A1 range references, optionally prefixed with a sheet name.
_RANGE = re.compile(
    r"^(?:(?:'((?:[^']|'')+)'|([A-Za-z_][\w. ]*))!)?"
    r"(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?$",
    re.IGNORECASE,
)
MAX_LIST_CHOICES = 5000


def _choices(formula, workbook, worksheet):
    """Return supported allowed values, or None when unsafe to infer."""
    if not formula or not isinstance(formula, str):
        return None
    val = formula.strip()
    if val.startswith("="):
        val = val[1:].strip()
    if val.startswith('"') and val.endswith('"'):
        # Standard Excel literal lists; quoted commas within individual
        # values and complex expressions are outside current scope.
        choices = [x.strip() for x in val[1:-1].split(",")]
        return choices if len(choices) <= MAX_LIST_CHOICES else None
    match = _RANGE.fullmatch(val)
    if not match:
        return None
    tab_name = (match.group(1) or match.group(2) or worksheet.title).replace("''", "'")
    if tab_name not in workbook.sheetnames:
        return None
    start = match.group(3)
    end = match.group(4) or start
    left, top, right, bottom = range_boundaries(start + ":" + end)
    if left != right or (bottom - top + 1) > MAX_LIST_CHOICES:
        return None
    tab = workbook[tab_name]
    choices = []
    for row in tab.iter_rows(min_col=left, max_col=right, min_row=top, max_row=bottom):
        cell = row[0]
        if cell.data_type == "f":
            return None  # Never evaluate formula-based allowed values.
        if cell.value is not None:
            choices.append(cell.value)
    return choices


def _equivalent(actual, option):
    if actual == option:
        return True
    text_actual = str(actual).strip()
    text_option = str(option).strip()
    if text_actual == text_option:
        return True
    try:
        return Decimal(text_actual) == Decimal(text_option)
    except (InvalidOperation, TypeError, ValueError):
        return False


def inspect_excel(data: bytes, df, is_xlsm: bool, worksheet_name=None):
    """Inspect one chosen worksheet and validate supported dropdown lists."""
    macro_embedded = False
    with ZipFile(io.BytesIO(data)) as archive:
        macro_embedded = any(
            part.lower().endswith("/vbaproject.bin") for part in archive.namelist()
        )

    book = load_workbook(io.BytesIO(data), read_only=False, data_only=False,
                         keep_vba=False, keep_links=False)
    try:
        worksheet = book[worksheet_name] if worksheet_name is not None else book.worksheets[0]
        formula_cells = [
            cell for cell in worksheet._cells.values()
            if cell.row >= 2 and cell.data_type == "f"
        ]
        formulas_present = bool(formula_cells)
        broken_formula_refs = sum(
            "#REF!" in str(cell.value) for cell in formula_cells
        )
        validations = list(worksheet.data_validations.dataValidation) if worksheet.data_validations else []
        listed = [rule for rule in validations if rule.type == "list"]
        validated_cells = 0
        invalid_cells = 0
        unsupported_rules = 0
        masks = {}
        # Formula source is inspected statically: #REF! is a broken A1
        # reference even when the workbook has no cached formula value.
        for cell in formula_cells:
            if (2 <= cell.row <= len(df) + 1 and
                    1 <= cell.column <= len(df.columns) and
                    "#REF!" in str(cell.value)):
                formula_mask = masks.setdefault(
                    cell.column - 1, np.zeros(len(df), dtype=bool)
                )
                formula_mask[cell.row - 2] = True
        rule_details = []
        for rule in listed:
            choices = _choices(rule.formula1, book, worksheet)
            rule_details.append({
                "ranges": str(rule.sqref)[:200],
                "source": "supported literal/range" if choices is not None else "unchecked",
                "allowed_values": [str(x)[:80] for x in choices[:100]] if choices else [],
                "allowed_values_truncated": bool(choices and len(choices) > 100),
            })
            if choices is None:
                unsupported_rules += 1
                continue
            # Skip blank cells; missing data is handled by the core profiler.
            for cell_range in rule.sqref.ranges:
                left, top, right, bottom = cell_range.bounds
                for column_index in range(max(1, left), min(right, len(df.columns)) + 1):
                    col = column_index - 1
                    mask = masks.setdefault(col, np.zeros(len(df), dtype=bool))
                    for excel_row in range(max(2, top), min(bottom, len(df) + 1) + 1):
                        pos = excel_row - 2
                        value = df.iat[pos, col]
                        if value is None or bool(pd.isna(value)):
                            continue
                        validated_cells += 1
                        if not any(_equivalent(value, choice) for choice in choices):
                            if not mask[pos]:
                                invalid_cells += 1
                            mask[pos] = True
        warnings = []
        if is_xlsm:
            warnings.append(
                "Macro-enabled Excel file read in analysis-only mode. "
                "VBA code is never executed, edited, or included in HTML exports."
            )
        if formulas_present:
            warnings.append(
                "Formula cells were detected. Data profiling uses saved/cached "
                "values; formulas are not recalculated, so missing or stale "
                "cached results may affect findings."
            )
        if broken_formula_refs:
            warnings.append(
                f"{broken_formula_refs} formula(s) contain an explicit #REF! reference."
            )
        if unsupported_rules:
            warnings.append(
                f"{unsupported_rules} dropdown/list rule(s) use dynamic, named, "
                "formula, or other unsupported sources and were not checked."
            )
        if validations and len(listed) != len(validations):
            warnings.append(
                "Non-list Excel data validations (such as dates or numeric ranges) "
                "are not interpreted as dropdown checks."
            )
        return {
            "masks": masks,
            "metadata": {
                "macro_enabled_format": is_xlsm,
                "vba_project_present": macro_embedded,
                "vba_executed": False,
                "worksheet_analyzed": worksheet.title,
                "sheets_in_workbook": len(book.sheetnames),
                "dropdown_rules": len(listed),
                "dropdown_rules_unchecked": unsupported_rules,
                "dropdown_cells_checked": validated_cells,
                "dropdown_invalid_cells": invalid_cells,
                "formula_cells_detected": formulas_present,
                "formula_cells_count": len(formula_cells),
                "broken_formula_refs": broken_formula_refs,
                "excel_tables": len(worksheet.tables),
                "table_names": list(worksheet.tables.keys())[:60],
                "conditional_formatting_ranges": len(worksheet.conditional_formatting),
                "worksheet_visibility": worksheet.sheet_state,
                "dropdown_rules_detail": rule_details[:80],
                "warnings": warnings,
            },
        }
    finally:
        book.close()
