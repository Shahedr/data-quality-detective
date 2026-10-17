"""Multi-sheet workbook overview, selection, complete export and VBA triage."""

import io
import sys
import types
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from webapp.app import app
from webapp.vba import audit_vba

client = TestClient(app)


def workbook_bytes(macro=False):
    book = Workbook()
    first = book.active
    first.title = "Orders"
    first.append(["order_id", "order_date", "status"])
    first.append([1, "2026-10-01", "Open"])
    first.append([2, "not-a-date", "Closed"])
    first.append([2, "not-a-date", "Closed"])
    first["D1"] = "formula"
    first["D2"] = "=A2+#REF!"
    dv = DataValidation(type="list", formula1='"Open,Closed"')
    first.add_data_validation(dv)
    dv.add("C2:C4")

    inventory = book.create_sheet("Inventory")
    inventory.append(["sku", "count"])
    inventory.append(["SKU-001", 4])
    inventory.append(["SKU-002", 5])
    inventory.append(["SKU-003", 5])

    hidden = book.create_sheet("Private Lookups")
    hidden.sheet_state = "hidden"
    hidden.append(["allowed"])
    hidden.append(["North"])
    hidden.append(["South"])

    output = io.BytesIO()
    book.save(output)
    book.close()
    if macro:
        output.seek(0, io.SEEK_END)
        with ZipFile(output, "a", ZIP_DEFLATED) as archive:
            archive.writestr("xl/vbaProject.bin", b"fake-vba-marker-for-testing-only")
    return output.getvalue()


def call(path, filename, data, fields=None):
    return client.post(path, data=fields or {}, files={"file": (filename, data)})


def test_all_sheets_listed_and_explicit_selection_changes_report():
    raw = workbook_bytes()
    res = call("/api/analyze", "company.xlsx", raw)
    assert res.status_code == 200, res.text[:500]
    body = res.json()
    assert body["selected_sheet"] == "Orders"
    assert [x["name"] for x in body["worksheets"]] == [
        "Orders", "Inventory", "Private Lookups"
    ]
    assert body["workbook_total_rows"] == 3 + 3 + 2
    assert body["summary"]["rows"] == 3
    assert body["workbook"]["formula_cells_count"] == 1
    assert body["workbook"]["broken_formula_refs"] == 1
    assert body["worksheets"][2]["visibility"] == "hidden"
    assert body["worksheets"][1]["rows"] == 3
    assert body["vba_audit"]["source_scan_status"] == "no_vba"

    inventory = call("/api/analyze", "company.xlsx", raw, {"sheet": "Inventory"})
    assert inventory.status_code == 200, inventory.text[:500]
    data = inventory.json()
    assert data["selected_sheet"] == "Inventory"
    assert data["summary"]["rows"] == 3
    assert data["summary"]["columns"] == 2
    assert data["spreadsheet"]["columns"] == ["sku", "count"]
    assert data["spreadsheet"]["rows"][0]["values"][0] == "SKU-001"

    hidden = call("/api/analyze", "company.xlsx", raw, {"sheet": "Private Lookups"})
    assert hidden.status_code == 200, hidden.text[:500]
    assert hidden.json()["spreadsheet"]["columns"] == ["allowed"]
    assert hidden.json()["workbook"]["worksheet_visibility"] == "hidden"


def test_virtual_grid_pages_and_issue_filters_are_sheet_specific():
    raw = workbook_bytes()
    inventory = call("/api/rows", "company.xlsx", raw, {
        "sheet": "Inventory", "issue": "all", "page": "1"
    })
    assert inventory.status_code == 200, inventory.text
    payload = inventory.json()
    assert payload["selected_sheet"] == "Inventory"
    assert payload["total_rows"] == 3
    assert payload["rows"][0]["values"][0] == "SKU-001"

    dupes = call("/api/rows", "company.xlsx", raw, {
        "sheet": "Orders", "issue": "duplicate", "page": "1"
    })
    assert dupes.status_code == 200, dupes.text
    assert [row["row_number"] for row in dupes.json()["rows"]] == [2, 3]

    no_duplicates = call("/api/rows", "company.xlsx", raw, {
        "sheet": "Inventory", "issue": "duplicate", "page": "1"
    })
    assert no_duplicates.status_code == 200
    assert no_duplicates.json()["rows"] == []
    assert call("/api/rows", "company.xlsx", raw, {"sheet": "Wrong"}).status_code == 422


def test_complete_html_export_contains_all_sheets_and_hidden_warning():
    raw = workbook_bytes()
    response = call("/api/export/html", "company.xlsx", raw)
    assert response.status_code == 200, response.text[:500]
    html = response.text
    assert "Workbook overview — all worksheets" in html
    assert "Worksheet: Orders" in html
    assert "Worksheet: Inventory" in html
    assert "Worksheet: Private Lookups" in html
    assert "hidden" in html
    assert "SKU-003" in html
    assert "not-a-date" in html
    assert "formula" in html
    assert html.count('<td class="row-index">') == 8
    assert "VBA code executed: No" in html
    assert '<script src=' not in html


def test_combined_workbook_rows_bound_applies_across_tabs(monkeypatch):
    from webapp import workbook
    monkeypatch.setattr(workbook, "MAX_ROWS", 6)
    response = call("/api/analyze", "company.xlsx", workbook_bytes())
    assert response.status_code == 413
    assert "across a workbook" in response.json()["detail"]


def test_macro_project_never_executes_and_reports_scanner_availability():
    raw = workbook_bytes(macro=True)
    response = call("/api/analyze", "test.xlsm", raw)
    assert response.status_code == 200, response.text[:500]
    macro = response.json()["vba_audit"]
    assert macro["vba_project_present"] is True
    assert macro["vba_executed"] is False
    assert macro["source_scan_status"] in {
        "not_installed", "scan_failed", "unreadable_or_empty_project"
    }


def test_vba_analyzer_extracts_modules_and_flags_with_fake_parser(monkeypatch):
    class FakeParser:
        def __init__(self, name, data):
            self.name, self.data = name, data
        def detect_vba_macros(self):
            return True
        def extract_macros(self):
            return [("embedded", "VBA/ThisWorkbook", "ThisWorkbook",
                     "Sub Workbook_Open()\n Shell ""cmd""\nEnd Sub")]
        def analyze_macros(self):
            return [("AutoExec", "Workbook_Open", "Runs at workbook open"),
                    ("Suspicious", "Shell", "Launches a process")]
        def close(self):
            pass

    parent = types.ModuleType("oletools")
    module = types.ModuleType("oletools.olevba")
    module.VBA_Parser = FakeParser
    monkeypatch.setitem(sys.modules, "oletools", parent)
    monkeypatch.setitem(sys.modules, "oletools.olevba", module)
    result = audit_vba("fake.xlsm", b"unused", True)
    assert result["source_scan_status"] == "scanned"
    assert result["modules"][0]["name"] == "ThisWorkbook"
    assert result["finding_counts"]["autoexec"] == 1
    assert result["finding_counts"]["suspicious"] == 1
    assert result["vba_executed"] is False
    assert "source code" not in result


def test_upload_index_shows_worksheet_picker_and_macro_results():
    home = client.get("/")
    assert home.status_code == 200
    assert 'id="sheetSelect"' in home.text
    assert 'id="sheetOverview"' in home.text
    assert 'id="macroNotice"' in home.text
    script = client.get("/spreadsheet.js").text
    assert 'body.append("sheet", state.sheet)' in script
