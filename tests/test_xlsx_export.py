"""Open generated color-coded XLSX review files to verify contents and safety."""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl import Workbook

from webapp.app import app

client = TestClient(app)


def sample_xlsx():
    book = Workbook()
    a = book.active
    a.title = "Orders"
    a.append(["ID", "Status", "Calculation"])
    a.append([1, "Open", "=1+2"])
    a.append([1, "Open", "=1+2"])
    a.append([3, "Invalid", "=3+4"])
    a.append([4, "", "notice"])
    dv = DataValidation(type="list", formula1='"Open,Closed"')
    a.add_data_validation(dv)
    dv.add("B2:B5")
    b = book.create_sheet("Private Lookup")
    b.sheet_state = "hidden"
    b.append(["Key"])
    b.append(["A"])
    b.append(["B"])
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def export(filename, payload):
    return client.post(
        "/api/export/xlsx", files={"file": (filename, payload)}
    )


def test_xlsx_export_contains_all_sheets_summary_and_issue_register():
    reply = export("messy.xlsx", sample_xlsx())
    assert reply.status_code == 200, reply.text[:200]
    assert "audit_review.xlsx" in reply.headers["content-disposition"]
    assert reply.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    result = load_workbook(BytesIO(reply.content), data_only=False)
    try:
        assert result.sheetnames == [
            "DQD Audit Overview", "DQD Issue Register", "Orders", "Private Lookup"
        ]
        a = result["Orders"]
        assert a.max_row == 5
        assert a.max_column == 5  # row number + duplicate + 3 original columns
        assert a["C2"].value == "1"
        assert a["D4"].value == "Invalid"
        assert a["D4"].fill.fgColor.rgb.endswith("EDDBFD")
        assert a["B2"].value == "Duplicate group"
        assert a["A3"].value == "2"
        # openpyxl-created formulas have no cached values. The values-only
        # export leaves them blank rather than fabricating a calculated result.
        assert a["E2"].value is None
        assert a["E2"].data_type != "f"
        assert result["Private Lookup"]["C2"].value == "A"
        issues = list(result["DQD Issue Register"].values)
        assert any(item[4] == "Value not in supported dropdown list"
                   and item[2] == "B4" for item in issues[1:])
        assert any(item[4] == "Duplicate row" for item in issues[1:])
        overview = result["DQD Audit Overview"]
        assert any(row[0] == "Orders" and row[5] == "1"
                   for row in list(overview.values))
        assert all(not cell.data_type == "f" for sheet in result
                   for row in sheet for cell in row)
    finally:
        result.close()


def test_csv_excel_formula_injection_is_saved_as_literal_text():
    csv = (
        'id,comment\n'
        '1,"=HYPERLINK(""https://bad.example"",""CLICK"")"\n'
        '2,"+SUM(1,2)"\n'
        '3,"@cmd"\n'
    ).encode()
    response = export("unsafe.csv", csv)
    assert response.status_code == 200, response.text[:200]
    workbook = load_workbook(BytesIO(response.content))
    try:
        sheet = workbook["CSV data"]
        assert sheet["D2"].value.startswith("=HYPERLINK")
        assert sheet["D2"].data_type == "s"
        assert sheet["D3"].value == "+SUM(1,2)"
        assert sheet["D3"].data_type == "s"
        assert sheet["D4"].value == "@cmd"
        assert sheet["D4"].data_type == "s"
        assert not workbook._external_links
    finally:
        workbook.close()


def test_audit_workbook_does_not_embed_vba_or_excel_formulas():
    data = sample_xlsx()
    response = export("source.xlsm", data)
    assert response.status_code == 200
    from zipfile import ZipFile
    with ZipFile(BytesIO(response.content)) as z:
        assert not any("vbaProject" in path for path in z.namelist())
    result = load_workbook(BytesIO(response.content))
    try:
        assert result["DQD Audit Overview"]["B3"].value.startswith("Value-only")
        assert result["Orders"]["E2"].value is None
        assert result["Orders"]["E2"].data_type != "f"
    finally:
        result.close()


def test_large_xlsx_export_returns_useful_413_not_crashes(monkeypatch):
    from webapp import xlsx_export
    monkeypatch.setattr(xlsx_export, "MAX_EXPORT_CELLS", 5)
    reply = export("messy.xlsx", sample_xlsx())
    assert reply.status_code == 413
    assert "Download full HTML spreadsheet" in reply.json()["detail"]


def test_generated_titles_do_not_collide_with_source_titles():
    book = Workbook()
    book.active.title = "DQD Audit Overview"
    book.active.append(["id"])
    book.active.append([1])
    book.create_sheet("DQD Issue Register").append(["id"])
    buf = BytesIO()
    book.save(buf)
    book.close()
    response = export("collision.xlsx", buf.getvalue())
    assert response.status_code == 200, response.text[:200]
    result = load_workbook(BytesIO(response.content))
    try:
        assert len(result.sheetnames) == 4
        assert "DQD Audit Overview" in result.sheetnames
        assert "DQD Issue Register" in result.sheetnames
        assert "DQD Audit Overview (2)" in result.sheetnames
        assert "DQD Issue Register (2)" in result.sheetnames
    finally:
        result.close()
