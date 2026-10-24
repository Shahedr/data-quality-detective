"""Optional sample-only mode and the real free hosted upload contract."""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from webapp.app import app


def test_hosted_site_keeps_local_large_file_setup_visible(monkeypatch):
    monkeypatch.setenv('VERCEL', '1')
    page = TestClient(app).get('/').text
    assert 'id="localSetup"' in page
    assert 'Process up to 1 GB locally' in page
    assert 'Windows &mdash; setup' in page
    assert 'Mac &mdash; Terminal setup' in page
    assert 'href="#localInstall"' in page
    assert 'href="http://127.0.0.1:8000/#largeFiles"' not in page
    assert 'Scripts\\python.exe -m webapp.launch_local' in page
    assert './.venv/bin/python -m webapp.launch_local' in page


def test_local_app_behaviour_and_public_demo_routes_are_separate(monkeypatch):
    monkeypatch.delenv("DQD_PUBLIC_DEMO_ONLY", raising=False)
    client = TestClient(app)
    config = client.get("/api/config")
    assert config.status_code == 200
    assert config.json()["mode"] == "local"
    assert config.json()["upload_allowed"] is True
    assert config.json()["macro_execution_via_web"] is False
    assert client.get("/api/demo/analyze").status_code == 404
    local = client.post("/api/analyze", files={
        "file": ("tiny.csv", b"a,b\n1,2\n")
    })
    assert local.status_code == 200, local.text


def test_public_demo_accepts_no_uploaded_workbooks(monkeypatch):
    monkeypatch.setenv("DQD_PUBLIC_DEMO_ONLY", "1")
    client = TestClient(app)
    config = client.get("/api/config").json()
    assert config["mode"] == "public_demo"
    assert config["upload_allowed"] is False
    assert config["macro_execution_via_web"] is False
    for endpoint in ("/api/analyze", "/api/rows", "/api/export/html", "/api/export/xlsx"):
        response = client.post(
            endpoint,
            files={"file": ("evil.xlsm", b"PKfake-evil-macro-body")},
            data={"page": "1", "issue": "all"},
            follow_redirects=False,
        )
        assert response.status_code == 403, (endpoint, response.text)
        assert "unavailable" in response.json()["detail"].lower()
        assert response.headers.get("X-Frame-Options") == "DENY"
    assert client.get("/api/demo/analyze", params={"file": "private.xlsm"}).status_code == 200


def test_demo_data_and_grid_filters(monkeypatch):
    monkeypatch.setenv("DQD_PUBLIC_DEMO_ONLY", "1")
    client = TestClient(app)
    analyzed = client.get("/api/demo/analyze")
    assert analyzed.status_code == 200, analyzed.text[:500]
    payload = analyzed.json()
    assert payload["summary"]["rows"] == 12
    assert payload["selected_sheet"] is None
    assert payload["workbook"] is None
    assert payload["vba_audit"] is None
    rows = client.get("/api/demo/rows", params={"page": 1, "issue": "all"})
    assert rows.status_code == 200
    assert rows.json()["total_rows"] == 12
    duplicate = client.get("/api/demo/rows", params={"page": 1, "issue": "duplicate"})
    assert duplicate.status_code == 200
    assert duplicate.json()["total_matching_rows"] == 2
    assert [r["row_number"] for r in duplicate.json()["rows"]] == [11, 12]
    assert client.get("/api/demo/rows", params={"page": 0}).status_code == 422
    assert client.get("/api/demo/rows", params={"issue": "bad"}).status_code == 422


def test_public_demo_reports_are_complete_and_safe(monkeypatch):
    monkeypatch.setenv("DQD_PUBLIC_DEMO_ONLY", "1")
    client = TestClient(app)
    html = client.get("/api/demo/export/html")
    assert html.status_code == 200, html.text[:300]
    assert "attachment" in html.headers["content-disposition"]
    assert html.text.count('<td class="row-index">') == 12
    report = client.get("/api/demo/export/xlsx")
    assert report.status_code == 200, report.text[:150]
    assert "sample_audit_review.xlsx" in report.headers["content-disposition"]
    book = load_workbook(BytesIO(report.content))
    try:
        assert "CSV data" in book.sheetnames
        assert "DQD Issue Register" in book.sheetnames
        assert book["CSV data"].max_row == 13
        assert all(
            cell.data_type != "f" for sheet in book for row in sheet for cell in row
        )
    finally:
        book.close()


def test_deployed_ui_includes_mode_detection_and_demo_mode_paging(monkeypatch):
    monkeypatch.setenv("DQD_PUBLIC_DEMO_ONLY", "1")
    client = TestClient(app)
    homepage = client.get("/")
    script = client.get("/spreadsheet.js")
    main_script = client.get("/app.js")
    assert homepage.status_code == 200
    assert main_script.status_code == 200
    assert '<script src="/app.js"></script>' in homepage.text
    assert "initializeMode();" in main_script.text
    assert "/api/demo/analyze" in main_script.text
    assert "/api/demo/export/xlsx" in main_script.text
    assert "/api/demo/rows" in script.text
    assert "state.demo" in script.text


def test_vercel_hosts_real_free_uploads_with_explicit_body_limit(monkeypatch):
    monkeypatch.delenv("DQD_PUBLIC_DEMO_ONLY", raising=False)
    monkeypatch.setenv("VERCEL", "1")
    client = TestClient(app)
    config = client.get("/api/config").json()
    assert config["mode"] == "free_web"
    assert config["upload_allowed"] is True
    assert config["max_upload_bytes"] == 4 * 1024 * 1024
    response = client.post("/api/analyze", files={
        "file": ("actual.csv", b"a,b\n1,2\n1,2\n")
    })
    assert response.status_code == 200, response.text
    assert response.json()["summary"]["duplicate_rows"] == 1
    assert client.get("/api/saas/config").status_code == 404
    assert client.get("/app").status_code == 404
    assert client.post(
        "/api/analyze",
        headers={"Content-Length": str(5 * 1024 * 1024)},
        content=b"x",
    ).status_code == 413


def test_retired_paid_routes_return_not_found(monkeypatch):
    monkeypatch.delenv("DQD_PUBLIC_DEMO_ONLY", raising=False)
    monkeypatch.setenv("VERCEL", "1")
    client = TestClient(app)
    for route in ("/app", "/product",
                  "/api/saas/config", "/api/saas/billing/checkout",
                  "/api/saas/analyze", "/api/saas/admin/config"):
        assert client.get(route).status_code == 404, route
