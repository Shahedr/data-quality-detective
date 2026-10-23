"""Real API flows and round-trip safety checks for the isolated workspace MVP."""
import hashlib
import io
import json
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from webapp.app import app

client = TestClient(app)
CSV = b'id,name,date,email\n001, Alice ,12/31/2025,a@example.com\n001, Alice ,12/31/2025,a@example.com\n002,BOB,not-a-date,bad\n003,NA,,\n004,=1+1,,\n'


def post(path, data=CSV, options=None, **fields):
    if options is not None:
        fields["options"] = json.dumps(options)
    return client.post(path, files={"file": ("input.csv", data, "text/csv")}, data=fields)


def test_clean_preview_confirm_and_literal_export_preserves_original():
    original_hash = hashlib.sha256(CSV).hexdigest()
    options = {"columns": ["name"], "trim": True, "case": "lower", "remove_duplicates": True}
    preview = post("/api/clean/preview", options=options)
    assert preview.status_code == 200, preview.text
    info = preview.json()
    assert info["duplicates_removed"] == 1
    assert info['copy_preview']['columns'][:2] == ['id', 'name']
    assert info['copy_preview']['rows'][0][:2] == ['001', 'alice']
    assert len(info['copy_preview']['rows']) == 4
    assert info["rows_before"] == 5 and info["rows_after"] == 4
    assert info["changes"][0] == {"row": 2, "column": "name", "before": " Alice ", "after": "alice"}
    exported = post("/api/clean/export", options=options, fingerprint=info["fingerprint"], confirmed="true")
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content), data_only=False)
    values = list(book.active.values)
    assert values[1][:2] == ("001", "alice")
    assert values[3][1] == "na"  # literal NA was never silently turned into missing
    assert values[4][1] == "=1+1"
    assert book.active["B5"].data_type == "s"
    assert hashlib.sha256(CSV).hexdigest() == original_hash
    book.close()


def test_strict_date_cleaning_preserves_invalid_values():
    options = {"columns": ["date"], "date_format": "%m/%d/%Y"}
    info = post("/api/clean/preview", options=options).json()
    assert info["cells_changed"] == 2
    assert info["invalid_dates_unchanged"] == 1
    assert info["changes"][0]["after"] == "2025-12-31"


@pytest.mark.parametrize("options", [{"columns": ["unknown"], "trim": True}, {"columns": ["name", "name"]}, {"case": "title"}, {"trim": "true"}, {"extra": 1}, {"trim": True}, {"date_format": "%x"}])
def test_reject_invalid_cleaning_options(options):
    assert post("/api/clean/preview", options=options).status_code == 422


def test_download_rejects_stale_file_options_and_missing_confirmation():
    options = {"columns": ["name"], "trim": True}
    info = post("/api/clean/preview", options=options).json()
    assert post("/api/clean/export", options=options, fingerprint=info["fingerprint"]).status_code == 409
    assert post("/api/clean/export", data=CSV + b'005,new,,\n', options=options, fingerprint=info["fingerprint"], confirmed="true").status_code == 409
    assert post("/api/clean/export", options={"columns": ["name"], "case": "upper"}, fingerprint=info["fingerprint"], confirmed="true").status_code == 409


def test_workbook_health_detects_all_sheets_errors_links_and_macro_parts():
    book = Workbook()
    tab = book.active
    tab.title = "Orders"
    tab.append(["id", "formula", "error"])
    tab.append([1, "=#REF!+1", "#DIV/0!"])
    hidden = book.create_sheet("Hidden data")
    hidden.sheet_state = "hidden"
    hidden.append(["code"])
    hidden.append(["001"])
    source = io.BytesIO()
    book.save(source)
    augmented = io.BytesIO()
    with ZipFile(source) as src, ZipFile(augmented, "w", ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
        dst.writestr("xl/vbaProject.bin", b"not executable")
    raw = augmented.getvalue()
    response = client.post("/api/workspace", files={"file": ("book.xlsm", raw)}, data={"sheet": "Orders"})
    assert response.status_code == 200, response.text
    health = response.json()["health"]
    assert health["formulas"] == 1 and health["broken_references"] == 1
    assert health["saved_errors"] == 1 and health["external_link_parts"] == 1
    assert health["macro_present"] is True
    assert health["sheets"][1]["visibility"] == "hidden"
    options = {"sheet": "Orders", "columns": ["id"]}
    preview = client.post("/api/clean/preview", files={"file": ("book.xlsm", raw)}, data={"options": json.dumps(options)}).json()
    export = client.post("/api/clean/export", files={"file": ("book.xlsm", raw)}, data={"options": json.dumps(options), "confirmed": "true", "fingerprint": preview["fingerprint"]})
    assert export.status_code == 200, export.text
    result = load_workbook(io.BytesIO(export.content))
    assert result.sheetnames == ["Orders", "Hidden data"]
    assert result["Hidden data"]["A2"].value == "001"
    assert result["Hidden data"].sheet_state == "hidden"
    assert result["Orders"]["B2"].value == "=#REF!+1"
    assert export.content == raw  # no approved edits means an exact original copy
    result.close()


def test_excel_literal_edits_preserve_every_other_package_part():
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "formula"])
    sheet.append([" Alice ", "=A2"])
    sheet["A2"].number_format = "@"
    hidden = book.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    original = io.BytesIO()
    book.save(original)
    packaged = io.BytesIO()
    with ZipFile(io.BytesIO(original.getvalue())) as src, ZipFile(packaged,"w",ZIP_DEFLATED) as dst:
        for part in src.infolist():
            dst.writestr(part,src.read(part.filename))
        dst.writestr("xl/vbaProject.bin", b"test macro bytes - never executed")
        dst.writestr("customXml/item1.xml", b"<test>retain this extension</test>")
    data = packaged.getvalue()
    options = {"columns": ["name"], "trim": True}
    preview = client.post("/api/clean/preview", files={"file":("macro.xlsm",data)},data={"options":json.dumps(options)})
    assert preview.status_code == 200, preview.text
    output = client.post("/api/clean/export",files={"file":("macro.xlsm",data)},data={"options":json.dumps(options),"confirmed":"true","fingerprint":preview.json()["fingerprint"]})
    assert output.status_code == 200, output.text
    assert '.xlsm' in output.headers['content-disposition']
    with ZipFile(io.BytesIO(data)) as a, ZipFile(io.BytesIO(output.content)) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert a.read(name) == b.read(name), name
    result = load_workbook(io.BytesIO(output.content))
    assert result.active['A2'].value == 'Alice'
    assert result.active['A2'].number_format == '@'
    assert result.active['B2'].value == '=A2'
    assert result['Hidden'].sheet_state == 'hidden'
    blocked = client.post("/api/clean/preview",files={"file":("macro.xlsm",data)},data={"options":json.dumps({"remove_duplicates":True})})
    assert blocked.status_code == 422


def test_privacy_pages_and_preferences_asset():
    for path in ['/privacy','/terms','/cookies']:
        r=client.get(path)
        assert r.status_code==200 and 'shahed2338@gmail.com' in r.text
        assert 'Shahed' in r.text and 'Public beta' in r.text
    script=client.get('/preferences.js').text
    assert 'localStorage.' not in script and 'document.cookie' not in script
    assert 'dark' in script and 'light' in script


def test_signed_workbook_and_formula_cached_text_cannot_be_changed():
    book=Workbook()
    book.active.append(['name'])
    book.active.append([' A '])
    original=io.BytesIO();book.save(original)
    def archive(extra, formula=False):
        result=io.BytesIO()
        with ZipFile(io.BytesIO(original.getvalue())) as src,ZipFile(result,'w',ZIP_DEFLATED) as dst:
            for part in src.infolist():
                data=src.read(part.filename)
                if formula and part.filename=='xl/worksheets/sheet1.xml':
                    data=data.replace(b'<c r="A2" t="inlineStr"><is><t xml:space="preserve"> A </t></is></c>',b'<c r="A2" t="str"><f>" A "</f><v> A </v></c>')
                dst.writestr(part,data)
            if extra:dst.writestr('_xmlsignatures/sig1.xml',b'<Signature/>')
        return result.getvalue()
    for data in [archive(True),archive(False,True)]:
        response=client.post('/api/clean/preview',files={'file':('protected.xlsx',data)},data={'options':json.dumps({'columns':['name'],'trim':True})})
        assert response.status_code==422,response.text


@pytest.mark.parametrize("template,data,expected", [
    ("sales", b'order_id,order_date,quantity,unit_price\n1,2025-01-01,1,10\n1,no,0,-1\n', [2,1,1,1]),
    ("inventory", b'sku,quantity,reorder_level\nA,2,1\nB,-1,no\n', [0,1,1]),
    ("customers", b'customer_id,email,name\n1,a@example.com,A\n2,bad,\n', [0,1,1]),
])
def test_templates_charts_and_written_report(template, data, expected):
    result = post("/api/workspace", data=data, template=template)
    assert result.status_code == 200, result.text
    payload = result.json()
    assert [check["count"] for check in payload["template_checks"]] == expected
    assert len(payload["charts"]) > 0
    assert "screening rules" in payload["narrative"]
    assert "Original files are unchanged" in payload["narrative"]
    assert payload["health"]["available"] is False


def test_unknown_template_missing_columns_and_sample_only_gates(monkeypatch):
    assert post("/api/workspace", template="unknown").status_code == 422
    result = post("/api/workspace", template="inventory").json()
    assert all(c["status"] == "not_checked" for c in result["template_checks"])
    monkeypatch.setenv("DQD_PUBLIC_DEMO_ONLY", "1")
    monkeypatch.delenv("VERCEL", raising=False)
    for path in ["/api/workspace", "/api/clean/preview", "/api/clean/export"]:
        assert post(path).status_code == 403


def test_new_routes_keep_hosted_upload_caps(monkeypatch):
    monkeypatch.setenv("VERCEL", "1")
    monkeypatch.delenv("DQD_FREE_UPLOADS_DISABLED", raising=False)
    for path in ["/api/workspace", "/api/clean/preview", "/api/clean/export"]:
        assert client.post(path, content=b'x' * (4 * 1024 * 1024 + 65537)).status_code == 413


def test_explorer_handles_blank_constant_and_infinite_numeric_columns():
    result = post("/api/workspace", data=b'a,b,c\n,1,inf\n,1,-inf\n,1,2\n')
    assert result.status_code == 200, result.text
    data = result.json()
    assert data["charts"][0]["missing"] == 3
    assert sum(x["count"] for x in data["charts"][1]["distribution"]) == 3


def test_extreme_charts_and_numeric_excel_headers():
    import pandas as pd
    from webapp.workspace import explore
    charts = explore(pd.DataFrame({"extreme": [-1.7e308, 1.7e308], "constant": [1.7e308, 1.7e308]}))
    assert sum(b["count"] for b in charts[0]["distribution"]) == 2
    assert charts[1]["distribution"] == [{"label": "1.7e+308", "count": 2}]
    book = Workbook()
    book.active.append([123])
    book.active.append([" A "])
    source = io.BytesIO()
    book.save(source)
    result = client.post("/api/clean/preview", files={"file": ("numeric.xlsx", source.getvalue())}, data={"options": json.dumps({"columns": ["123"], "trim": True})})
    assert result.status_code == 200, result.text
    assert result.json()["changes"][0]["after"] == "A"


def test_clean_export_rejects_unrepresentable_text_without_truncating():
    data = ('name\n' + 'x' * 32768 + '\n').encode()
    preview = post("/api/clean/preview", data=data).json()
    response = post("/api/clean/export", data=data, fingerprint=preview["fingerprint"], confirmed="true")
    assert response.status_code == 422
    assert "cannot be represented safely" in response.text


def test_change_preview_is_bounded_and_large_dataset_explorer():
    data = ('sku,quantity,reorder_level\n' + ''.join(f' A{i} ,{i},2\n' for i in range(1001))).encode()
    result = post("/api/clean/preview", data=data, options={"columns": ["sku"], "trim": True}).json()
    assert result["cells_changed"] == 1001 and len(result["changes"]) == 100
    assert result["preview_truncated"] is True
    report = post("/api/workspace", data=data, template="inventory")
    assert report.status_code == 200
    assert len(report.json()["charts"][1]["distribution"]) <= 8


def test_template_identifiers_preserve_leading_zeros_and_literal_na():
    result = post("/api/workspace", data=b'customer_id,email,name\n001,a@example.com,NA\n1,b@example.com,B\n', template="customers")
    assert result.status_code == 200
    assert [c["count"] for c in result.json()["template_checks"]] == [0,0,0]
