import io
import json
import struct
import re
import asyncio
import time
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from webapp.app import app
from webapp import largefiles
from webapp.largeclean import clean_excel
from webapp.workspace import CleanOptions
from webapp.zipguard import check_directory
from starlette.requests import Request


@pytest.fixture(autouse=True)
def jobs(monkeypatch):
    monkeypatch.delenv('VERCEL', raising=False)
    monkeypatch.setattr(largefiles, '_jobs', {})
    yield
    for job in largefiles._jobs.values():
        largefiles.remove_files(job)


def upload(client, name, data, clean):
    response = client.post('/api/large/uploads', json={'name': name, 'size': len(data), 'clean': clean})
    assert response.status_code == 200, response.text
    job = response.json()
    endpoint = '/api/large/uploads/'+job['id']
    token = {'X-Upload-Token': job['token']}
    assert client.put(endpoint, content=data, headers={**token, 'X-Upload-Offset': '0'}).status_code == 200
    assert client.post(endpoint+'/complete', headers=token).status_code == 200
    return job, endpoint, token, client.get(endpoint, headers=token).json()


def test_csv_full_clean_preview_confirmation_stream_download_and_cleanup():
    client = TestClient(app)
    data = b'id,name,date\n001, Alice ,31/12/2025\n001, Alice ,31/12/2025\n002,BOB,invalid\n'
    job, endpoint, token, result = upload(client, 'test.csv', data, {'columns': ['name', 'date'], 'trim': True,
                       'case': 'lower', 'date_format': '%d/%m/%Y', 'remove_duplicates': True})
    assert result['status'] == 'review', result
    preview = result['result']['cleaning']
    assert preview['duplicates_removed'] == 1 and preview['rows_after'] == 2
    assert preview['cells_changed'] == 5
    assert not largefiles._jobs[job['id']]['path'].exists()
    assert client.post(endpoint+'/download', data={'token': job['token'], 'fingerprint': 'wrong', 'confirmed': 'true'}).status_code == 409
    assert client.post(endpoint+'/download', data={'token': job['token'], 'fingerprint': preview['fingerprint']}).status_code == 409
    response = client.post(endpoint+'/download', data={'token': job['token'], 'fingerprint': preview['fingerprint'], 'confirmed': 'true'})
    assert response.status_code == 200, response.text
    assert response.content == b'id,name,date\r\n001,alice,2025-12-31\r\n002,bob,invalid\r\n'
    assert not largefiles._jobs[job['id']]['output'].exists()
    assert client.post(endpoint+'/download', data={'token': job['token'], 'fingerprint': preview['fingerprint'], 'confirmed': 'true'}).status_code == 409


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Orders'
    sheet.append(['name', 'calculation'])
    sheet.append([' Alice ', '=1+2'])
    sheet['A2'].number_format = '@'
    book.create_sheet('Hidden').sheet_state = 'hidden'
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def test_large_excel_literal_cleaning_and_all_other_parts_preserved():
    original = workbook()
    client = TestClient(app)
    job, endpoint, token, result = upload(client, 'test.xlsm', original, {'sheet': 'Orders', 'columns': ['A', 'B'], 'trim': True})
    assert result['status'] == 'review', result
    preview = result['result']['cleaning']
    assert preview['cells_changed'] == 1
    response = client.post(endpoint+'/download', data={'token': job['token'], 'fingerprint': preview['fingerprint'], 'confirmed': 'true'})
    with ZipFile(io.BytesIO(original)) as source, ZipFile(io.BytesIO(response.content)) as result:
        assert source.namelist() == result.namelist()
        for name in source.namelist():
            if name != 'xl/worksheets/sheet1.xml':
                assert source.read(name) == result.read(name), name
    book = load_workbook(io.BytesIO(response.content), data_only=False)
    assert book['Orders']['A2'].value == 'Alice'
    assert book['Orders']['A2'].number_format == '@'
    assert book['Orders']['B2'].value == '=1+2'
    assert book['Hidden'].sheet_state == 'hidden'
    book.close()


def test_indented_inline_string_does_not_gain_newlines(tmp_path):
    source, output = tmp_path/'in.xlsx', tmp_path/'out.xlsx'
    with ZipFile(io.BytesIO(workbook())) as original, ZipFile(source, 'w') as patched:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'<is><t xml:space="preserve"> Alice </t></is>', b'<is>\n<t>Alice</t>\n</is>')
            patched.writestr(item, data)
    clean_excel(source, output, CleanOptions(sheet='Orders', columns=['A'], case='upper'), lambda: None)
    book = load_workbook(output)
    assert book['Orders']['A2'].value == 'ALICE'
    book.close()
    assert not output.with_suffix('.sqlite').exists()


def test_prefixed_cells_are_patched_without_matching_worksheet_prefix(tmp_path):
    source, output = tmp_path/'in.xlsx', tmp_path/'out.xlsx'
    with ZipFile(io.BytesIO(workbook())) as original, ZipFile(source, 'w') as patched:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'xmlns=', b'xmlns:c=')
                data = re.sub(rb'(<\/?)([A-Za-z][A-Za-z0-9]*)', rb'\1c:\2', data)
            patched.writestr(item, data)
    clean_excel(source, output, CleanOptions(sheet='Orders', columns=['A'], trim=True), lambda: None)
    book = load_workbook(output)
    assert book['Orders']['A2'].value == 'Alice'
    assert book['Orders']['B2'].value == '=1+2'
    book.close()


@pytest.mark.parametrize('options', [
    {'columns': ['unknown'], 'trim': True},
    {'columns': ['name', 'name'], 'trim': True},
    {'columns': [], 'trim': True},
])
def test_invalid_csv_cleaning_leaves_no_temporary_files(options):
    job, _, _, result = upload(TestClient(app), 'bad.csv', b'name\nA\n', options)
    assert result['status'] == 'failed'
    assert not largefiles._jobs[job['id']]['output'].exists()
    assert not largefiles._jobs[job['id']]['path'].exists()


def test_sparse_coordinates_rejected_before_pandas(monkeypatch):
    from webapp import workbook as module
    book = Workbook()
    book.active['A1'] = 'header'
    book.active['A100002'] = 'last'
    data = io.BytesIO()
    book.save(data)
    book.close()
    def forbidden(*args, **kwargs):
        raise AssertionError('pandas should not be invoked')
    monkeypatch.setattr(module.pd, 'ExcelFile', forbidden)
    response = TestClient(app).post('/api/analyze', files={'file': ('sparse.xlsx', data.getvalue())})
    assert response.status_code in {413, 422}


def test_zip_directory_is_bounded_before_constructor():
    data = io.BytesIO()
    with ZipFile(data, 'w') as archive:
        for index in range(50):
            archive.writestr(str(index), b'')
    with pytest.raises(ValueError):
        check_directory(io.BytesIO(data.getvalue()), 10, 1024**2)


def test_standard_zip64_directory_rejected_even_without_legacy_sentinels(monkeypatch):
    import zipfile
    monkeypatch.setattr(zipfile, 'ZIP_FILECOUNT_LIMIT', 1)
    data = io.BytesIO()
    with ZipFile(data, 'w') as archive:
        archive.writestr('one', b'')
        archive.writestr('two', b'')
    with pytest.raises(ValueError, match='ZIP64'):
        check_directory(io.BytesIO(data.getvalue()), 100, 1024**2)


def test_large_formatted_empty_row_rejected_before_pandas(monkeypatch):
    from webapp import workbook as module
    book = Workbook()
    book.active['A1'] = 'header'
    book.active.row_dimensions[100002].height = 15
    data = io.BytesIO()
    book.save(data)
    book.close()
    def forbidden(*args, **kwargs):
        raise AssertionError('pandas should not be invoked')
    monkeypatch.setattr(module.pd, 'ExcelFile', forbidden)
    assert TestClient(app).post('/api/analyze', files={'file': ('formatted.xlsx', data.getvalue())}).status_code in {413, 422}


def test_failed_download_headers_still_remove_output():
    client = TestClient(app)
    job, endpoint, token, result = upload(client, 'test.csv', b'name\n A \n', {'columns': ['name'], 'trim': True})
    scope = {'type': 'http', 'method': 'POST', 'scheme': 'http', 'path': endpoint+'/download',
             'server': ('testserver', 80), 'client': ('testclient', 1), 'headers': [], 'asgi': {'spec_version': '2.4'}}
    response = largefiles.download_cleaned(job['id'], Request(scope), token=job['token'],
                    fingerprint=result['result']['cleaning']['fingerprint'], confirmed=True)
    async def unavailable(*args):
        raise RuntimeError('Connection closed before headers')
    with pytest.raises(RuntimeError):
        asyncio.run(response(scope, unavailable, unavailable))
    current = largefiles._jobs[job['id']]
    assert current['status'] == 'complete' and not current['output'].exists()


def test_download_not_started_expires_and_releases_job():
    client = TestClient(app)
    job, _, _, _ = upload(client, 'test.csv', b'name\n A \n', {'columns': ['name'], 'trim': True})
    current = largefiles._jobs[job['id']]
    current.update(status='downloading', download_active=False, download_requested=time.time()-61)
    largefiles.cleanup()
    assert job['id'] not in largefiles._jobs and not current['output'].exists()


def test_validation_cache_respects_changed_resource_limits(monkeypatch):
    from webapp import workbook as module
    module._validated.clear()
    data = workbook()
    assert module.validate_upload('first.xlsx', data) == '.xlsx'
    assert module.validate_upload('same.xlsx', data) == '.xlsx'
    monkeypatch.setattr(module, 'MAX_ROWS', 0)
    with pytest.raises(Exception):
        module.validate_upload('same.xlsx', data)


def test_native_download_origin_policy_keeps_external_requests_blocked():
    client = TestClient(app)
    assert client.get('/').headers['referrer-policy'] == 'same-origin'
    assert client.post('/api/large/uploads', json={'name': 'a.csv', 'size': 10}, headers={'Origin': 'null'}).status_code == 403


def test_csv_cell_limit_is_checked_before_pandas(monkeypatch):
    from webapp import workbook as module
    monkeypatch.setattr(module, 'MAX_CELLS', 6)
    def forbidden(*args, **kwargs):
        raise AssertionError('pandas must not allocate an oversized CSV')
    monkeypatch.setattr(module.pd, 'read_csv', forbidden)
    result = TestClient(app).post('/api/analyze', files={'file': ('wide.csv', b'a,b\n1,2\n3,4\n5,6\n')})
    assert result.status_code == 413, result.text


def test_csv_preflight_counts_quoted_whitespace_and_widest_row(monkeypatch):
    from webapp import workbook as module
    monkeypatch.setattr(module, 'MAX_CELLS', 5)
    for data in [b'a\n" "\n" "\n" "\n" "\n" "\n', b'a\n1,2,3\n4\n']:
        with pytest.raises(Exception) as error:
            module.validate_upload('bounded.csv', data)
        assert error.value.status_code == 413
