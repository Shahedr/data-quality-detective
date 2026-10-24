"""Packaged local app: a small control window and the existing browser interface."""
import json
import os
import sys
import threading
import time
import webbrowser
from urllib.request import urlopen, Request

from webapp.launch_local import available_port


class LocalApp:
    def __init__(self):
        # A desktop download must always use the local request boundaries.
        os.environ.pop('VERCEL', None)
        os.environ.pop('DQD_PUBLIC_DEMO_ONLY', None)
        import uvicorn
        from webapp.app import app
        self.url = f'http://127.0.0.1:{available_port()}'
        self.error = None
        self.server = uvicorn.Server(uvicorn.Config(
            app, host='127.0.0.1', port=int(self.url.rsplit(':', 1)[1]),
            proxy_headers=False, log_config=None, access_log=False,
            timeout_graceful_shutdown=10,
        ))
        self.thread = threading.Thread(target=self.run, daemon=True)

    def run(self):
        try:
            self.server.run()
        except BaseException as exc:
            self.error = str(exc) or 'The local service stopped.'

    def stop(self):
        self.server.should_exit = True
        self.thread.join(timeout=15)
        return not self.thread.is_alive()


def self_test():
    import tkinter as tk
    root = tk.Tk()
    root.withdraw()
    root.update()
    root.destroy()
    local = LocalApp()
    local.thread.start()
    try:
        for _ in range(300):
            if local.server.started:
                break
            if local.error:
                raise RuntimeError(local.error)
            time.sleep(.1)
        else:
            raise RuntimeError('Packaged app did not start')
        with urlopen(local.url + '/api/config', timeout=10) as response:
            config = json.loads(response.read())
        assert config['mode'] == 'local' and config['large_upload_bytes'] == 1024**3
        with urlopen(local.url, timeout=10) as response:
            assert b'id="localSetup"' in response.read()
        boundary = 'DesktopSmokeBoundary'
        payload = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="test.csv"\r\n'
                   'Content-Type: text/csv\r\n\r\nid,value\n1,10\n2,20\n'
                   f'\r\n--{boundary}--\r\n').encode()
        request = Request(local.url + '/api/analyze', data=payload,
                          headers={'Content-Type': 'multipart/form-data; boundary=' + boundary})
        with urlopen(request, timeout=30) as response:
            assert json.loads(response.read())['summary']['rows'] == 2
        def api(path, method='GET', data=None, headers=None):
            with urlopen(Request(local.url+path, data=data, method=method, headers=headers or {}), timeout=30) as response:
                return json.loads(response.read())
        sample = b'id,value\n1,10\n2,20\n'
        job = api('/api/large/uploads', 'POST', json.dumps({'name': 'grid.csv', 'size': len(sample)}).encode(), {'Content-Type': 'application/json'})
        path = '/api/large/uploads/'+job['id']
        token = {'X-Upload-Token': job['token']}
        api(path, 'PUT', sample, {**token, 'X-Upload-Offset': '0'})
        api(path+'/complete', 'POST', headers=token)
        for _ in range(100):
            state = api(path, headers=token)
            if state['status'] not in {'queued', 'processing'}:
                break
            time.sleep(.1)
        assert state['status'] == 'complete', state
        grid = api(path+'/rows', headers=token)
        assert grid['total'] == 3 and grid['rows'][2]['values'] == ['2', '20']
        api(path, 'DELETE', headers=token)
        # Exercise the same Excel/cleaning routes inside each platform's frozen runtime.
        from io import BytesIO
        from openpyxl import Workbook, load_workbook
        from urllib.parse import urlencode
        book = Workbook()
        book.active.title = 'Orders'
        book.active.append(['Customer', 'Formula'])
        for _ in range(65):
            book.active.append([' Alice ', '=2+3'])
        book.create_sheet('Inventory').append(['Widget', 42])
        book.create_sheet('Hidden').append(['Private reference'])
        book['Hidden'].sheet_state = 'hidden'
        source = BytesIO()
        book.save(source)
        book.close()
        job = api('/api/large/uploads', 'POST', json.dumps({
            'name': 'workbook.xlsx', 'size': len(source.getvalue()),
            'clean': {'sheet': 'Orders', 'columns': ['A'], 'trim': True},
        }).encode(), {'Content-Type': 'application/json'})
        path, token = '/api/large/uploads/'+job['id'], {'X-Upload-Token': job['token']}
        api(path, 'PUT', source.getvalue(), {**token, 'X-Upload-Offset': '0'})
        api(path+'/complete', 'POST', headers=token)
        for _ in range(100):
            state = api(path, headers=token)
            if state['status'] not in {'queued', 'processing'}:
                break
            time.sleep(.1)
        assert state['status'] == 'review', state
        original = api(path+'/rows?sheet=Orders&offset=50', headers=token)
        cleaned = api(path+'/rows?sheet=Orders&offset=50&cleaned=true', headers=token)
        assert original['rows'][0]['values'] == [' Alice ', '=2+3']
        assert cleaned['rows'][0]['values'] == ['Alice', '=2+3']
        assert api(path+'/rows?sheet=Inventory', headers=token)['rows'][0]['values'] == ['Widget', '42']
        assert api(path+'/rows?sheet=Hidden', headers=token)['rows'][0]['values'] == ['Private reference']
        data = urlencode({'token': job['token'], 'fingerprint': state['result']['cleaning']['fingerprint'], 'confirmed': 'true'}).encode()
        with urlopen(Request(local.url+path+'/download', data=data, headers={'Content-Type': 'application/x-www-form-urlencoded'}), timeout=30) as response:
            downloaded = load_workbook(BytesIO(response.read()), data_only=False)
        assert downloaded['Orders']['A2'].value == 'Alice' and downloaded['Orders']['B2'].value == '=2+3'
        assert downloaded['Hidden'].sheet_state == 'hidden'
        downloaded.close()
        api(path, 'DELETE', headers=token)
        if '--large-file-test' in sys.argv:
            import tempfile
            from pathlib import Path
            size = 1024**3
            with tempfile.TemporaryDirectory(prefix='dqd-packaged-fixture-') as directory:
                fixture = Path(directory) / 'one-gib.csv'
                block = (b'x'*1023+b'\n')*4096
                with fixture.open('wb') as target:
                    target.write(b'value\n')
                    remaining = size-6
                    while remaining:
                        portion = block[:min(remaining, len(block))]
                        target.write(portion)
                        remaining -= len(portion)
                assert fixture.stat().st_size == size
                job = api('/api/large/uploads', 'POST', json.dumps({'name': fixture.name, 'size': size}).encode(), {'Content-Type': 'application/json'})
                path, token = '/api/large/uploads/'+job['id'], {'X-Upload-Token': job['token']}
                with fixture.open('rb') as source:
                    offset = 0
                    for chunk in iter(lambda: source.read(job['chunk_bytes']), b''):
                        api(path, 'PUT', chunk, {**token, 'X-Upload-Offset': str(offset)})
                        offset += len(chunk)
                api(path+'/complete', 'POST', headers=token)
                for _ in range(3000):
                    state = api(path, headers=token)
                    if state['status'] not in {'queued', 'processing'}:
                        break
                    time.sleep(.1)
                assert state['status'] == 'complete', state
                assert state['result']['bytes'] == size and state['result']['rows'] == 1048576
                final = api(path+'/rows?offset=1048576', headers=token)
                assert final['total'] == 1048577 and final['rows'][0]['row_number'] == 1048577
                assert final['rows'][0]['values'] == ['x'*500]
                api(path, 'DELETE', headers=token)
    finally:
        assert local.stop(), 'Packaged service did not stop cleanly'
    return 0


def main():
    if '--self-test' in sys.argv:
        try:
            return self_test()
        except Exception:
            import traceback
            from pathlib import Path
            Path('desktop-self-test.log').write_text(traceback.format_exc(), encoding='utf-8')
            return 1
    import tkinter as tk
    from tkinter import ttk
    root = tk.Tk()
    root.title('Data Quality Detective')
    root.geometry('490x210')
    root.resizable(False, False)
    panel = ttk.Frame(root, padding=24)
    panel.pack(fill='both', expand=True)
    ttk.Label(panel, text='Data Quality Detective', font=('', 17, 'bold')).pack(anchor='w')
    status = tk.StringVar(value='Starting your local workspace...')
    ttk.Label(panel, textvariable=status, wraplength=435).pack(anchor='w', pady=16)
    local = None
    opened = False
    def open_app():
        if local and local.server.started:
            webbrowser.open(local.url)
    button = ttk.Button(panel, text='Open workspace', command=open_app, state='disabled')
    button.pack(side='left')
    def close():
        status.set('Closing the local workspace...')
        root.update_idletasks()
        if local:
            local.stop()
        root.destroy()
    ttk.Button(panel, text='Quit', command=close).pack(side='right')
    root.protocol('WM_DELETE_WINDOW', close)
    try:
        local = LocalApp()
        local.thread.start()
    except Exception as exc:
        status.set('Could not start the app: ' + str(exc))
    def check():
        nonlocal opened
        if local and local.error:
            status.set('Could not start the app: ' + local.error)
        elif local and local.server.started:
            status.set('Your workspace is ready. Keep this window open while you work.\n' + local.url)
            button.state(['!disabled'])
            if not opened:
                opened = True
                open_app()
        root.after(250, check)
    root.after(100, check)
    root.mainloop()
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
